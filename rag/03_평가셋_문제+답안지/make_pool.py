"""
2단계: 새 답안지 초안 만들기 (pool_snap1016.xlsx)

질문마다 후보 공고를 모으고, 점수(score)를 미리 채울 수 있는 건 채움
  후보 = V0 상위 10 + V1 상위 10 + 1차 정답(1점 이상) 중 342건 안에 있는 것
  점수 미리 채우기
    ① 1차 답안지에 같은 (질문, 공고) 점수가 있으면 그대로 가져옴      → 할일 "1차점수"
    ② 타지역 공고 / 사용자와 다른 구 전용 공고 → 0점 (1차 채점 기준)  → 할일 "자동0"
    ③ 나머지는 빈칸                                                   → 할일 "채점필요"
"""
import sys
from pathlib import Path
import pandas as pd

HERE = Path(__file__).resolve().parent
RAG = HERE.parent
sys.path.insert(0, str(RAG / "04_실험_시험+채점결과"))
from retrievers import V0, V1, DOC_BY_ID

QS  = pd.read_csv(HERE / "search_questions.csv", encoding="utf-8-sig", dtype=str).fillna("")
OLD = pd.read_csv(RAG / "_보관_1차_snap1013" / "pool_labeled.csv", encoding="utf-8-sig", dtype=str).fillna("")
OUT = HERE / "pool_snap1016.xlsx"

old_score = {(r.qid, r.pblancId): (int(r.score), r["AI_이유"]) for _, r in OLD.iterrows()}

print("검색기 준비 중 (bge-m3 첫 실행은 몇 분 걸림)...")
v0, v1 = V0(), V1()

rows = []
for _, q in QS.iterrows():
    r0, r1 = v0.search(q.question, 10), v1.search(q.question, 10)
    old_gold = [pid for (qid, pid), (s, _) in old_score.items()
                if qid == q.qid and s >= 1 and pid in DOC_BY_ID]
    for pid in dict.fromkeys(r0 + r1 + old_gold):           # 합치고 중복 제거 (순서 유지)
        d = DOC_BY_ID[pid]
        src = "+".join(n for n, r in [("V0", r0), ("V1", r1)] if pid in r) or "1차정답"

        if (q.qid, pid) in old_score:
            score, why, todo = old_score[(q.qid, pid)][0], old_score[(q.qid, pid)][1], "1차점수"
        elif d["region"] == "타지역":
            score, why, todo = 0, "타지역 공고", "자동0"
        elif d["seoul_gu"] and q.region and q.region not in d["seoul_gu"]:
            score, why, todo = 0, f'{"·".join(d["seoul_gu"])} 전용 (사용자 {q.region})', "자동0"
        else:
            score, why, todo = None, "", "채점필요"

        rows.append({
            "할일": todo, "qid": q.qid, "질문": q.question, "유형": q.difficulty,
            "필요지원": q.need, "사용자지역": q.region,
            "공고명": d["pblancNm"], "공고지역": d["region"], "서울구": " ".join(d["seoul_gu"]),
            "지원대상": d["trgetNm"], "분야": f'{d["field_L"]}/{d["field_M"]}',
            "사업개요": d["summary"][:300],
            "score": score, "점수이유": why, "메모": "",
            "출처": src,
            "V0순위": r0.index(pid) + 1 if pid in r0 else None,
            "V1순위": r1.index(pid) + 1 if pid in r1 else None,
            "pblancId": pid, "URL": d["url"],
        })

pool = pd.DataFrame(rows)

# ── 엑셀로 저장 (보기 좋게) ─────────────────────────────────
from openpyxl.styles import PatternFill, Alignment, Font
COLOR = {"채점필요": "FFF2CC", "1차점수": "E2EFDA", "자동0": "EDEDED"}
WIDTH = {"할일": 9, "qid": 5, "질문": 30, "유형": 4, "필요지원": 9, "사용자지역": 8, "공고명": 45,
         "공고지역": 7, "서울구": 8, "지원대상": 8, "분야": 14, "사업개요": 60, "score": 6,
         "점수이유": 30, "메모": 20, "출처": 8, "V0순위": 6, "V1순위": 6, "pblancId": 22, "URL": 20}

with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    pool.to_excel(xw, sheet_name="답안지", index=False)
    ws = xw.sheets["답안지"]
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(pool.columns, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = WIDTH.get(col, 10)
        ws.cell(1, i).font = Font(bold=True)
    score_col = list(pool.columns).index("score") + 1
    for r in range(2, len(pool) + 2):
        fill = PatternFill("solid", fgColor=COLOR[ws.cell(r, 1).value])
        for c in range(1, len(pool.columns) + 1):
            ws.cell(r, c).fill = fill
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=c in (3, 7, 12))
        ws.cell(r, score_col).font = Font(bold=True)

    guide = pd.DataFrame({"점수": [2, 1, 0], "뜻": [
        "사용자가 원하는 지원 그 자체 (뭐 받을 수 있냐는 질문엔: 통합공고 또는 업종·상황에 딱 맞춘 지원)",
        "원하는 것과 비슷하거나 더 넓은 범위 (소상공인이면 누구나 받는 일반 지원)",
        "무관함, 또는 지역·특정 대상·산업이 달라서 못 씀"]})
    guide.to_excel(xw, sheet_name="채점기준", index=False)
    xw.sheets["채점기준"].column_dimensions["B"].width = 90

# ── 요약 ────────────────────────────────────────────────────
print(f"\n저장: {OUT}")
print("전체 후보:", len(pool), "줄 / 질문당 평균", round(len(pool) / len(QS), 1))
print(pool["할일"].value_counts().to_dict())
print("\n질문별 채점필요 개수:")
print(pool[pool["할일"] == "채점필요"].groupby("qid").size().to_string())
